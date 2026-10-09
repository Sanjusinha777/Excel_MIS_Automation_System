"""
Professional Excel MIS Report Generator
Uses openpyxl to generate styled, multi-tab executive workbooks with formulas,
KPI scorecards, category summaries, monthly trends, and data audits.
"""

import io
from datetime import datetime
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class ExcelMISReporter:
    def __init__(self, engine):
        self.engine = engine
        
        # Color Palette - Corporate Navy & Steel
        self.C_NAVY_DARK = "1B365D"       # Deep Navy Header
        self.C_NAVY_MED = "2E75B6"        # Medium Blue Subheader
        self.C_NAVY_LIGHT = "D9E1F2"      # Soft Blue Total / Highlight
        self.C_CARD_BG = "F2F4F8"         # Soft Gray/Blue KPI card background
        self.C_WHITE = "FFFFFF"
        self.C_GRAY_LIGHT = "F9FBFD"      # Zebra alternate
        self.C_BORDER = "D9D9D9"          # Light border
        self.C_ACCENT_GREEN = "2E7D32"    # Positive accent
        
        # Standard Borders
        self.thin_side = Side(border_style="thin", color=self.C_BORDER)
        self.cell_border = Border(top=self.thin_side, left=self.thin_side, right=self.thin_side, bottom=self.thin_side)
        
        self.double_bottom = Side(border_style="double", color="1B365D")
        self.top_thin = Side(border_style="thin", color="1B365D")
        self.total_border = Border(top=self.top_thin, bottom=self.double_bottom, left=self.thin_side, right=self.thin_side)
        
    def generate_workbook(self) -> openpyxl.Workbook:
        """Construct full styled Excel MIS workbook."""
        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)
        
        self._build_executive_summary_sheet(wb)
        self._build_category_analysis_sheet(wb)
        
        if self.engine.primary_date:
            self._build_monthly_trends_sheet(wb)
            
        self._build_data_quality_sheet(wb)
        self._build_cleaned_data_sheet(wb)
        
        return wb

    def save_to_file(self, filepath: str):
        """Save report workbook to disk."""
        wb = self.generate_workbook()
        wb.save(filepath)
        return filepath

    def save_to_bytes(self) -> bytes:
        """Save report workbook to in-memory bytes for direct web download."""
        wb = self.generate_workbook()
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    def _apply_sheet_view_defaults(self, ws):
        """Ensure grid lines are always visible."""
        ws.views.sheetView[0].showGridLines = True

    def _auto_fit_columns(self, ws, max_len_cap=45):
        """Auto fit column widths with padding."""
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in col:
                # Skip merged or title banner rows that span across
                if cell.row in [1, 2, 3] and cell.value and len(str(cell.value)) > 30:
                    continue
                if cell.value is not None:
                    cell_len = len(str(cell.value))
                    if cell_len > max_len:
                        max_len = cell_len
            col_width = min(max(max_len + 4, 12), max_len_cap)
            ws.column_dimensions[col_letter].width = col_width

    def _build_executive_summary_sheet(self, wb):
        ws = wb.create_sheet(title="Executive_Summary")
        self._apply_sheet_view_defaults(ws)
        
        # 1. Main Title Banner
        ws.merge_cells("A1:H2")
        title_cell = ws["A1"]
        title_cell.value = "EXECUTIVE MANAGEMENT INFORMATION SYSTEM (MIS) REPORT"
        title_cell.font = Font(name="Calibri", size=16, bold=True, color=self.C_WHITE)
        title_cell.fill = PatternFill(start_color=self.C_NAVY_DARK, end_color=self.C_NAVY_DARK, fill_type="solid")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Subtitle metadata
        ws.merge_cells("A3:H3")
        sub_cell = ws["A3"]
        gen_time = datetime.now().strftime("%d-%b-%Y %I:%M %p")
        sheet_source = self.engine.sheet_name or "Sheet1"
        sub_cell.value = f"Generated On: {gen_time}  |  Source Sheet: {sheet_source}  |  Total Records Analyzed: {len(self.engine.cleaned_df):,}"
        sub_cell.font = Font(name="Calibri", size=10, italic=True, color="555555")
        sub_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # 2. KPI Cards Block
        kpis = self.engine.get_summary_kpis()
        metric_name = kpis.get("metric_name", "Total Records")
        tot_val = kpis.get("total_value", 0)
        avg_val = kpis.get("avg_value", 0)
        max_val = kpis.get("max_value", 0)
        cnt_val = kpis.get("total_records", 0)

        # 4 KPI cards arranged horizontally
        cards = [
            ("TOTAL TRANSACTIONS", f"{cnt_val:,}", "Total volume count"),
            (f"TOTAL {metric_name.upper()}", f"{tot_val:,.2f}", "Cumulative total"),
            (f"AVERAGE {metric_name.upper()}", f"{avg_val:,.2f}", "Per-record average"),
            (f"MAX {metric_name.upper()}", f"{max_val:,.2f}", "Single record peak")
        ]
        
        cols = [("A", "B"), ("C", "D"), ("E", "F"), ("G", "H")]
        row_top = 5
        row_val = 6
        row_sub = 7
        
        for idx, (title, val_str, subtext) in enumerate(cards):
            c_start, c_end = cols[idx]
            ws.merge_cells(f"{c_start}{row_top}:{c_end}{row_top}")
            ws.merge_cells(f"{c_start}{row_val}:{c_end}{row_val}")
            ws.merge_cells(f"{c_start}{row_sub}:{c_end}{row_sub}")
            
            # Header
            c1 = ws[f"{c_start}{row_top}"]
            c1.value = title
            c1.font = Font(name="Calibri", size=9, bold=True, color="555555")
            c1.fill = PatternFill(start_color=self.C_CARD_BG, end_color=self.C_CARD_BG, fill_type="solid")
            c1.alignment = Alignment(horizontal="center", vertical="center")
            
            # Value
            c2 = ws[f"{c_start}{row_val}"]
            c2.value = val_str
            c2.font = Font(name="Calibri", size=14, bold=True, color=self.C_NAVY_DARK)
            c2.fill = PatternFill(start_color=self.C_CARD_BG, end_color=self.C_CARD_BG, fill_type="solid")
            c2.alignment = Alignment(horizontal="center", vertical="center")
            
            # Subtext
            c3 = ws[f"{c_start}{row_sub}"]
            c3.value = subtext
            c3.font = Font(name="Calibri", size=8, italic=True, color="777777")
            c3.fill = PatternFill(start_color=self.C_CARD_BG, end_color=self.C_CARD_BG, fill_type="solid")
            c3.alignment = Alignment(horizontal="center", vertical="center")

            # Border outer for card
            for r in range(row_top, row_sub + 1):
                for col_letter in [c_start, c_end]:
                    ws[f"{col_letter}{r}"].border = self.cell_border
        
        # 3. Executive Insights Section
        ws.cell(row=9, column=1, value="EXECUTIVE KEY FINDINGS & STRATEGIC HIGHLIGHTS").font = Font(name="Calibri", size=12, bold=True, color=self.C_NAVY_DARK)
        
        insights = self.engine.generate_executive_insights()
        cur_row = 10
        for ins in insights:
            # Clean markdown bold asterisks for clean Excel display
            clean_text = ins.replace("**", "")
            ws.merge_cells(start_row=cur_row, start_column=1, end_row=cur_row, end_column=8)
            cell = ws.cell(row=cur_row, column=1, value=f"•  {clean_text}")
            cell.font = Font(name="Calibri", size=10, color="333333")
            cell.alignment = Alignment(vertical="center")
            cur_row += 1
            
        cur_row += 1 # Spacing
        
        # 4. Top Breakdown Snapshot Table (Primary Dimension)
        if self.engine.primary_dimension:
            dim_df = self.engine.get_dimension_breakdown(self.engine.primary_dimension, top_n=6)
            if len(dim_df) > 0:
                ws.cell(row=cur_row, column=1, value=f"KEY PERFORMANCE BY {self.engine.primary_dimension.upper()}").font = Font(name="Calibri", size=12, bold=True, color=self.C_NAVY_DARK)
                cur_row += 1
                
                # Table Headers
                headers = [self.engine.primary_dimension, "Total Amount", "Avg Ticket", "Count", "Share %", "Cumulative %"]
                for c_idx, h_text in enumerate(headers, start=1):
                    c = ws.cell(row=cur_row, column=c_idx, value=h_text)
                    c.font = Font(name="Calibri", size=10, bold=True, color=self.C_WHITE)
                    c.fill = PatternFill(start_color=self.C_NAVY_MED, end_color=self.C_NAVY_MED, fill_type="solid")
                    c.alignment = Alignment(horizontal="center" if c_idx > 1 else "left", vertical="center")
                    c.border = self.cell_border
                cur_row += 1
                
                # Rows
                table_start = cur_row
                for r_idx, row in dim_df.iterrows():
                    ws.cell(row=cur_row, column=1, value=str(row[self.engine.primary_dimension])).alignment = Alignment(horizontal="left")
                    c_tot = ws.cell(row=cur_row, column=2, value=float(row.get('Total', 0)))
                    c_tot.number_format = "#,##0.00"
                    c_avg = ws.cell(row=cur_row, column=3, value=float(row.get('Average', 0)))
                    c_avg.number_format = "#,##0.00"
                    c_cnt = ws.cell(row=cur_row, column=4, value=int(row.get('Count', 0)))
                    c_cnt.number_format = "#,##0"
                    c_pct = ws.cell(row=cur_row, column=5, value=float(row.get('Share_Pct', 0)) / 100)
                    c_pct.number_format = "0.0%"
                    c_cum = ws.cell(row=cur_row, column=6, value=float(row.get('Cumulative_Pct', 0)) / 100)
                    c_cum.number_format = "0.0%"
                    
                    # Zebra striping
                    fill_c = self.C_GRAY_LIGHT if r_idx % 2 == 1 else self.C_WHITE
                    for ci in range(1, 7):
                        cell = ws.cell(row=cur_row, column=ci)
                        cell.font = Font(name="Calibri", size=10)
                        cell.fill = PatternFill(start_color=fill_c, end_color=fill_c, fill_type="solid")
                        cell.border = self.cell_border
                    cur_row += 1

                # Totals Row
                tot_c1 = ws.cell(row=cur_row, column=1, value="Grand Total")
                tot_c1.font = Font(name="Calibri", size=10, bold=True)
                
                tot_val_c = ws.cell(row=cur_row, column=2, value=f"=SUM(B{table_start}:B{cur_row-1})")
                tot_val_c.number_format = "#,##0.00"
                tot_val_c.font = Font(name="Calibri", size=10, bold=True)
                
                tot_avg_c = ws.cell(row=cur_row, column=3, value=f"=AVERAGE(C{table_start}:C{cur_row-1})")
                tot_avg_c.number_format = "#,##0.00"
                tot_avg_c.font = Font(name="Calibri", size=10, bold=True)
                
                tot_cnt_c = ws.cell(row=cur_row, column=4, value=f"=SUM(D{table_start}:D{cur_row-1})")
                tot_cnt_c.number_format = "#,##0"
                tot_cnt_c.font = Font(name="Calibri", size=10, bold=True)
                
                tot_pct_c = ws.cell(row=cur_row, column=5, value=1.0)
                tot_pct_c.number_format = "0.0%"
                tot_pct_c.font = Font(name="Calibri", size=10, bold=True)
                
                tot_cum_c = ws.cell(row=cur_row, column=6, value=1.0)
                tot_cum_c.number_format = "0.0%"
                tot_cum_c.font = Font(name="Calibri", size=10, bold=True)
                
                for ci in range(1, 7):
                    c = ws.cell(row=cur_row, column=ci)
                    c.fill = PatternFill(start_color=self.C_NAVY_LIGHT, end_color=self.C_NAVY_LIGHT, fill_type="solid")
                    c.border = self.total_border
                
        self._auto_fit_columns(ws)

    def _build_category_analysis_sheet(self, wb):
        ws = wb.create_sheet(title="Category_Analysis")
        self._apply_sheet_view_defaults(ws)
        
        # Title
        ws.cell(row=1, column=1, value="MULTI-DIMENSIONAL CATEGORY BREAKDOWN & SHARE ANALYSIS").font = Font(name="Calibri", size=14, bold=True, color=self.C_NAVY_DARK)
        
        dims_to_show = self.engine.dimension_cols[:4] if self.engine.dimension_cols else []
        cur_row = 3
        
        for dim in dims_to_show:
            dim_df = self.engine.get_dimension_breakdown(dim, top_n=10)
            if len(dim_df) == 0:
                continue
                
            ws.cell(row=cur_row, column=1, value=f"Analysis by: {dim.upper()}").font = Font(name="Calibri", size=11, bold=True, color=self.C_NAVY_MED)
            cur_row += 1
            
            headers = [dim, "Total Volume/Value", "Average", "Count", "% Share", "Cumulative %"]
            for c_idx, h_text in enumerate(headers, start=1):
                c = ws.cell(row=cur_row, column=c_idx, value=h_text)
                c.font = Font(name="Calibri", size=10, bold=True, color=self.C_WHITE)
                c.fill = PatternFill(start_color=self.C_NAVY_MED, end_color=self.C_NAVY_MED, fill_type="solid")
                c.border = self.cell_border
            cur_row += 1
            
            table_start = cur_row
            for r_idx, row in dim_df.iterrows():
                ws.cell(row=cur_row, column=1, value=str(row[dim]))
                c_tot = ws.cell(row=cur_row, column=2, value=float(row.get('Total', 0)))
                c_tot.number_format = "#,##0.00"
                c_avg = ws.cell(row=cur_row, column=3, value=float(row.get('Average', 0)))
                c_avg.number_format = "#,##0.00"
                c_cnt = ws.cell(row=cur_row, column=4, value=int(row.get('Count', 0)))
                c_cnt.number_format = "#,##0"
                c_pct = ws.cell(row=cur_row, column=5, value=float(row.get('Share_Pct', 0)) / 100)
                c_pct.number_format = "0.0%"
                c_cum = ws.cell(row=cur_row, column=6, value=float(row.get('Cumulative_Pct', 0)) / 100)
                c_cum.number_format = "0.0%"
                
                fill_c = self.C_GRAY_LIGHT if r_idx % 2 == 1 else self.C_WHITE
                for ci in range(1, 7):
                    c = ws.cell(row=cur_row, column=ci)
                    c.font = Font(name="Calibri", size=10)
                    c.fill = PatternFill(start_color=fill_c, end_color=fill_c, fill_type="solid")
                    c.border = self.cell_border
                cur_row += 1
                
            # Totals
            ws.cell(row=cur_row, column=1, value="Total").font = Font(name="Calibri", size=10, bold=True)
            c_tot_sum = ws.cell(row=cur_row, column=2, value=f"=SUM(B{table_start}:B{cur_row-1})")
            c_tot_sum.number_format = "#,##0.00"
            c_tot_sum.font = Font(name="Calibri", size=10, bold=True)
            c_tot_avg = ws.cell(row=cur_row, column=3, value=f"=AVERAGE(C{table_start}:C{cur_row-1})")
            c_tot_avg.number_format = "#,##0.00"
            c_tot_avg.font = Font(name="Calibri", size=10, bold=True)
            c_tot_cnt = ws.cell(row=cur_row, column=4, value=f"=SUM(D{table_start}:D{cur_row-1})")
            c_tot_cnt.number_format = "#,##0"
            c_tot_cnt.font = Font(name="Calibri", size=10, bold=True)
            c_tot_pct = ws.cell(row=cur_row, column=5, value=1.0)
            c_tot_pct.number_format = "0.0%"
            c_tot_pct.font = Font(name="Calibri", size=10, bold=True)
            c_tot_cum = ws.cell(row=cur_row, column=6, value=1.0)
            c_tot_cum.number_format = "0.0%"
            c_tot_cum.font = Font(name="Calibri", size=10, bold=True)
            
            for ci in range(1, 7):
                c = ws.cell(row=cur_row, column=ci)
                c.fill = PatternFill(start_color=self.C_NAVY_LIGHT, end_color=self.C_NAVY_LIGHT, fill_type="solid")
                c.border = self.total_border
                
            cur_row += 2 # Spacing between dimension tables
            
        self._auto_fit_columns(ws)

    def _build_monthly_trends_sheet(self, wb):
        ws = wb.create_sheet(title="Trends_Monthly")
        self._apply_sheet_view_defaults(ws)
        
        ws.cell(row=1, column=1, value="TIME SERIES TREND & MONTH-OVER-MONTH (MoM) RUN-RATE").font = Font(name="Calibri", size=14, bold=True, color=self.C_NAVY_DARK)
        
        trend_df = self.engine.get_time_series_trend(freq="M")
        if len(trend_df) == 0:
            return
            
        cur_row = 3
        headers = ["Period (YYYY-MM)", "Total Volume/Value", "Count", "Average", "MoM Growth %"]
        for c_idx, h_text in enumerate(headers, start=1):
            c = ws.cell(row=cur_row, column=c_idx, value=h_text)
            c.font = Font(name="Calibri", size=10, bold=True, color=self.C_WHITE)
            c.fill = PatternFill(start_color=self.C_NAVY_MED, end_color=self.C_NAVY_MED, fill_type="solid")
            c.border = self.cell_border
        cur_row += 1
        
        table_start = cur_row
        for r_idx, row in trend_df.iterrows():
            ws.cell(row=cur_row, column=1, value=str(row['Period']))
            c_tot = ws.cell(row=cur_row, column=2, value=float(row.get('Total', 0)))
            c_tot.number_format = "#,##0.00"
            c_cnt = ws.cell(row=cur_row, column=3, value=int(row.get('Count', 0)))
            c_cnt.number_format = "#,##0"
            c_avg = ws.cell(row=cur_row, column=4, value=float(row.get('Avg', 0)))
            c_avg.number_format = "#,##0.00"
            
            growth_val = float(row.get('Growth_Pct', 0)) / 100
            c_gw = ws.cell(row=cur_row, column=5, value=growth_val)
            c_gw.number_format = "+0.0%;-0.0%;0.0%"
            
            fill_c = self.C_GRAY_LIGHT if r_idx % 2 == 1 else self.C_WHITE
            for ci in range(1, 6):
                c = ws.cell(row=cur_row, column=ci)
                c.font = Font(name="Calibri", size=10)
                c.fill = PatternFill(start_color=fill_c, end_color=fill_c, fill_type="solid")
                c.border = self.cell_border
            cur_row += 1
            
        # Summary row
        ws.cell(row=cur_row, column=1, value="Grand Total / Avg").font = Font(name="Calibri", size=10, bold=True)
        c_tot = ws.cell(row=cur_row, column=2, value=f"=SUM(B{table_start}:B{cur_row-1})")
        c_tot.number_format = "#,##0.00"
        c_tot.font = Font(name="Calibri", size=10, bold=True)
        c_cnt = ws.cell(row=cur_row, column=3, value=f"=SUM(C{table_start}:C{cur_row-1})")
        c_cnt.number_format = "#,##0"
        c_cnt.font = Font(name="Calibri", size=10, bold=True)
        c_avg = ws.cell(row=cur_row, column=4, value=f"=AVERAGE(D{table_start}:D{cur_row-1})")
        c_avg.number_format = "#,##0.00"
        c_avg.font = Font(name="Calibri", size=10, bold=True)
        c_gw = ws.cell(row=cur_row, column=5, value="-")
        c_gw.alignment = Alignment(horizontal="center")
        c_gw.font = Font(name="Calibri", size=10, bold=True)
        
        for ci in range(1, 6):
            c = ws.cell(row=cur_row, column=ci)
            c.fill = PatternFill(start_color=self.C_NAVY_LIGHT, end_color=self.C_NAVY_LIGHT, fill_type="solid")
            c.border = self.total_border
            
        self._auto_fit_columns(ws)

    def _build_data_quality_sheet(self, wb):
        ws = wb.create_sheet(title="Data_Quality_Audit")
        self._apply_sheet_view_defaults(ws)
        
        ws.cell(row=1, column=1, value="DATA INTEGRITY & EXCEPTION AUDIT REPORT").font = Font(name="Calibri", size=14, bold=True, color=self.C_NAVY_DARK)
        
        audit = self.engine.audit_data_quality()
        
        # Summary Overview
        ws.cell(row=3, column=1, value="AUDIT METRICS SUMMARY").font = Font(name="Calibri", size=11, bold=True, color=self.C_NAVY_MED)
        
        kpi_audit = [
            ("Total Rows", f"{audit.get('total_rows', 0):,}"),
            ("Total Columns", str(audit.get('total_columns', 0))),
            ("Dataset Completeness", f"{audit.get('completeness_pct', 100)}%"),
            ("Duplicate Rows", f"{audit.get('duplicate_rows', 0):,}"),
            ("Missing Cell Count", f"{audit.get('missing_cells', 0):,}")
        ]
        
        cur_row = 4
        for k, v in kpi_audit:
            ws.cell(row=cur_row, column=1, value=k).font = Font(name="Calibri", size=10, bold=True)
            ws.cell(row=cur_row, column=2, value=v).font = Font(name="Calibri", size=10)
            ws.cell(row=cur_row, column=1).border = self.cell_border
            ws.cell(row=cur_row, column=2).border = self.cell_border
            cur_row += 1
            
        cur_row += 2
        
        # Missing columns
        ws.cell(row=cur_row, column=1, value="COLUMNS WITH NULL / MISSING VALUES").font = Font(name="Calibri", size=11, bold=True, color=self.C_NAVY_MED)
        cur_row += 1
        
        headers = ["Column Name", "Missing Rows", "Missing %"]
        for c_idx, h in enumerate(headers, start=1):
            c = ws.cell(row=cur_row, column=c_idx, value=h)
            c.font = Font(name="Calibri", size=10, bold=True, color=self.C_WHITE)
            c.fill = PatternFill(start_color=self.C_NAVY_MED, end_color=self.C_NAVY_MED, fill_type="solid")
            c.border = self.cell_border
        cur_row += 1
        
        miss_cols = audit.get("missing_columns", [])
        if not miss_cols:
            ws.cell(row=cur_row, column=1, value="No missing values found in any column (100% clean).")
            cur_row += 1
        else:
            for row in miss_cols:
                ws.cell(row=cur_row, column=1, value=row["Column"])
                c2 = ws.cell(row=cur_row, column=2, value=row["Missing_Count"])
                c2.number_format = "#,##0"
                c3 = ws.cell(row=cur_row, column=3, value=row["Missing_Pct"] / 100)
                c3.number_format = "0.0%"
                for ci in range(1, 4):
                    ws.cell(row=cur_row, column=ci).border = self.cell_border
                cur_row += 1
                
        cur_row += 2
        
        # Outliers Table
        ws.cell(row=cur_row, column=1, value="STATISTICAL OUTLIERS DETECTED (IQR METHOD)").font = Font(name="Calibri", size=11, bold=True, color=self.C_NAVY_MED)
        cur_row += 1
        
        out_headers = ["Metric Column", "Outlier Count", "Outlier %", "Min Outlier", "Max Outlier", "Expected Range"]
        for c_idx, h in enumerate(out_headers, start=1):
            c = ws.cell(row=cur_row, column=c_idx, value=h)
            c.font = Font(name="Calibri", size=10, bold=True, color=self.C_WHITE)
            c.fill = PatternFill(start_color=self.C_NAVY_MED, end_color=self.C_NAVY_MED, fill_type="solid")
            c.border = self.cell_border
        cur_row += 1
        
        outliers = audit.get("outliers", [])
        if not outliers:
            ws.cell(row=cur_row, column=1, value="No extreme statistical outliers detected.")
        else:
            for o in outliers:
                ws.cell(row=cur_row, column=1, value=o["Metric"])
                c2 = ws.cell(row=cur_row, column=2, value=o["Outlier_Count"])
                c2.number_format = "#,##0"
                c3 = ws.cell(row=cur_row, column=3, value=o["Outlier_Pct"] / 100)
                c3.number_format = "0.0%"
                c4 = ws.cell(row=cur_row, column=4, value=o["Min_Outlier"])
                c4.number_format = "#,##0.00"
                c5 = ws.cell(row=cur_row, column=5, value=o["Max_Outlier"])
                c5.number_format = "#,##0.00"
                ws.cell(row=cur_row, column=6, value=o["Expected_Range"])
                for ci in range(1, 7):
                    ws.cell(row=cur_row, column=ci).border = self.cell_border
                cur_row += 1

        self._auto_fit_columns(ws)

    def _build_cleaned_data_sheet(self, wb):
        ws = wb.create_sheet(title="Cleaned_Data")
        self._apply_sheet_view_defaults(ws)
        
        df = self.engine.cleaned_df
        if df is None:
            return
            
        # Headers
        for col_idx, col_name in enumerate(df.columns, start=1):
            c = ws.cell(row=1, column=col_idx, value=col_name)
            c.font = Font(name="Calibri", size=10, bold=True, color=self.C_WHITE)
            c.fill = PatternFill(start_color=self.C_NAVY_DARK, end_color=self.C_NAVY_DARK, fill_type="solid")
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = self.cell_border

        # Enable Freeze Panes and AutoFilter
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}"
        
        # Write rows
        for row_idx, row in enumerate(df.itertuples(index=False), start=2):
            fill_c = self.C_GRAY_LIGHT if row_idx % 2 == 1 else self.C_WHITE
            for col_idx, val in enumerate(row, start=1):
                col_name = df.columns[col_idx - 1]
                cell = ws.cell(row=row_idx, column=col_idx)
                
                if pd.isna(val):
                    cell.value = ""
                elif isinstance(val, (datetime, pd.Timestamp)):
                    cell.value = val.strftime("%Y-%m-%d")
                elif isinstance(val, (int, float, pd.Int64Dtype)):
                    cell.value = val
                    if col_name in self.engine.metric_cols:
                        cell.number_format = "#,##0.00" if isinstance(val, float) else "#,##0"
                else:
                    cell.value = str(val)
                    
                cell.font = Font(name="Calibri", size=9)
                cell.fill = PatternFill(start_color=fill_c, end_color=fill_c, fill_type="solid")
                cell.border = self.cell_border
                
        self._auto_fit_columns(ws, max_len_cap=35)
